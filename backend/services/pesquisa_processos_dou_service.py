# ============================================================
# SERVICE — Pesquisa determinística de processos no DOU
#
# Regras:
# - sem IA;
# - sem similaridade;
# - sem correspondência parcial;
# - pesquisa exclusivamente números completos de processo;
# - aceita variações de pontuação/espaçamento;
# - gera resultado auditável por publicação e por processo.
# ============================================================

from __future__ import annotations

import csv
import datetime
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, Iterable, Sequence


# Mantido apenas por compatibilidade com chamadas antigas.
# A fonte operacional dos processos é a coluna PROCESSO da aba PROCESSOS
# do arquivo palavras_chave.xlsx.
PROCESSOS_PADRAO: tuple[str, ...] = ()

# Estrutura esperada:
# 5 dígitos + 6 dígitos + 4 dígitos + 2 dígitos.
#
# Os separadores podem ser ponto, barra, hífen ou espaço.
# O limite de 5 caracteres evita unir números muito distantes no texto.
PADRAO_CANDIDATO_PROCESSO = re.compile(
    r"(?<!\d)"
    r"(\d{5})[\s.\-/]{0,5}"
    r"(\d{6})[\s.\-/]{0,5}"
    r"(\d{4})[\s.\-/]{0,5}"
    r"(\d{2})"
    r"(?!\d)"
)

CAMPOS_PESQUISA_PADRAO: tuple[str, ...] = (
    "texto_integral",
    "texto",
    "titulo",
    "titulo_publicacao",
    "titulo_listagem",
    "ementa",
    "identifica",
    "sub_titulo",
    "resumo",
    "resumo_listagem",
)

COLUNAS_CSV: tuple[str, ...] = (
    "processo",
    "processo_normalizado",
    "status",
    "data_publicacao",
    "secao",
    "pagina",
    "jornal",
    "orgao",
    "titulo",
    "ementa",
    "identifica",
    "sub_titulo",
    "quantidade_ocorrencias_publicacao",
    "campos_encontrados",
    "campo_primeira_ocorrencia",
    "valor_encontrado",
    "trecho_encontrado",
    "url",
    "arquivo_xml",
    "zip",
    "arquivo_origem",
    "fonte",
    "origem",
    "id_publicacao",
    "hash_conteudo",
    "indice_base",
    "chave_resultado",
)


# ============================================================
# UTILITÁRIOS
# ============================================================

def agora_iso() -> str:
    return datetime.datetime.now().isoformat(timespec="seconds")


def texto_seguro(valor: Any) -> str:
    if valor is None:
        return ""
    return str(valor)


def normalizar_processo(valor: Any) -> str:
    """
    Mantém apenas os dígitos do processo.

    Exemplo:
    25351.413100/2023-48 -> 25351413100202348
    """
    return re.sub(r"\D", "", texto_seguro(valor))


def formatar_processo_normalizado(valor: str) -> str:
    """
    Formata uma sequência de 17 dígitos no padrão:
    00000.000000/0000-00
    """
    normalizado = normalizar_processo(valor)

    if len(normalizado) != 17:
        raise ValueError(
            f"Processo inválido: {valor!r}. "
            "O número normalizado deve possuir exatamente 17 dígitos."
        )

    return (
        f"{normalizado[0:5]}."
        f"{normalizado[5:11]}/"
        f"{normalizado[11:15]}-"
        f"{normalizado[15:17]}"
    )


def preparar_processos(
    processos: Sequence[str] | None = None,
) -> dict[str, str]:
    """
    Retorna:
        {
            "processo_normalizado": "processo_formatado"
        }
    """
    processos_entrada = (
        tuple(processos)
        if processos is not None
        else PROCESSOS_PADRAO
    )
    mapa: dict[str, str] = {}

    for processo in processos_entrada:
        normalizado = normalizar_processo(processo)

        if len(normalizado) != 17:
            raise ValueError(
                f"Processo inválido: {processo!r}. "
                "Após remover a pontuação, são esperados 17 dígitos."
            )

        formatado = formatar_processo_normalizado(normalizado)

        if normalizado in mapa:
            raise ValueError(
                f"Processo duplicado na configuração: {formatado}"
            )

        mapa[normalizado] = formatado

    if not mapa:
        raise ValueError("Nenhum processo foi informado para a pesquisa.")

    return mapa


def carregar_processos_planilha(
    caminho_planilha: Path | str,
    nome_aba: str = "PROCESSOS",
    nome_coluna: str = "PROCESSO",
) -> tuple[str, ...]:
    """
    Lê somente a coluna PROCESSO da aba PROCESSOS.

    Regras:
    - ignora células vazias;
    - aceita diferenças de maiúsculas/minúsculas nos nomes da aba/coluna;
    - valida exatamente 17 dígitos após remover pontuação;
    - formata todos os números em 00000.000000/0000-00;
    - elimina duplicidades preservando a ordem da planilha.
    """

    caminho = Path(caminho_planilha)

    if not caminho.exists():
        raise FileNotFoundError(
            f"Planilha de palavras-chave não encontrada: {caminho}"
        )

    try:
        from openpyxl import load_workbook
    except ImportError as erro:
        raise RuntimeError(
            "A leitura de palavras_chave.xlsx requer o pacote openpyxl."
        ) from erro

    workbook = load_workbook(
        filename=caminho,
        read_only=True,
        data_only=True,
    )

    try:
        mapa_abas = {
            str(nome).strip().casefold(): nome
            for nome in workbook.sheetnames
        }
        chave_aba = nome_aba.strip().casefold()

        if chave_aba not in mapa_abas:
            raise ValueError(
                f"Aba {nome_aba!r} não encontrada em {caminho}. "
                f"Abas disponíveis: {workbook.sheetnames}"
            )

        planilha = workbook[mapa_abas[chave_aba]]
        cabecalhos = next(
            planilha.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            ),
            (),
        )

        indice_coluna = None
        chave_coluna = nome_coluna.strip().casefold()

        for indice, valor in enumerate(cabecalhos):
            if str(valor or "").strip().casefold() == chave_coluna:
                indice_coluna = indice
                break

        if indice_coluna is None:
            raise ValueError(
                f"Coluna {nome_coluna!r} não encontrada na aba "
                f"{nome_aba!r} de {caminho}."
            )

        processos: list[str] = []
        vistos: set[str] = set()
        invalidos: list[str] = []

        for numero_linha, linha in enumerate(
            planilha.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            if indice_coluna >= len(linha):
                continue

            valor = linha[indice_coluna]

            if valor is None or not str(valor).strip():
                continue

            normalizado = normalizar_processo(valor)

            if len(normalizado) != 17:
                invalidos.append(
                    f"linha {numero_linha}: {valor!r} "
                    f"({len(normalizado)} dígitos)"
                )
                continue

            formatado = formatar_processo_normalizado(normalizado)

            if normalizado in vistos:
                continue

            vistos.add(normalizado)
            processos.append(formatado)

        if invalidos:
            detalhes = "; ".join(invalidos)
            raise ValueError(
                "Foram encontrados processos inválidos na coluna "
                f"{nome_coluna}: {detalhes}"
            )

        if not processos:
            raise ValueError(
                f"Nenhum processo preenchido na coluna {nome_coluna!r} "
                f"da aba {nome_aba!r}."
            )

        return tuple(processos)

    finally:
        workbook.close()


def montar_trecho(
    texto: str,
    inicio: int,
    fim: int,
    margem: int = 220,
) -> str:
    """
    Gera um trecho curto e auditável em torno da ocorrência.
    """
    texto = texto_seguro(texto)

    inicio_trecho = max(0, inicio - margem)
    fim_trecho = min(len(texto), fim + margem)

    trecho = texto[inicio_trecho:fim_trecho]
    trecho = re.sub(r"\s+", " ", trecho).strip()

    if inicio_trecho > 0:
        trecho = f"...{trecho}"

    if fim_trecho < len(texto):
        trecho = f"{trecho}..."

    return trecho


def primeiro_valor(
    registro: dict,
    chaves: Sequence[str],
    padrao: Any = "",
) -> Any:
    for chave in chaves:
        valor = registro.get(chave)
        if valor not in (None, "", [], {}):
            return valor
    return padrao


# ============================================================
# LOCALIZAÇÃO EXATA
# ============================================================

def localizar_processos_no_texto(
    texto: Any,
    processos_normalizados: set[str],
) -> list[dict]:
    """
    Extrai candidatos com estrutura completa de processo e mantém
    apenas os que correspondem exatamente à lista pesquisada.
    """
    texto_original = texto_seguro(texto)

    if not texto_original:
        return []

    ocorrencias: list[dict] = []

    for correspondencia in PADRAO_CANDIDATO_PROCESSO.finditer(texto_original):
        normalizado = "".join(correspondencia.groups())

        if normalizado not in processos_normalizados:
            continue

        ocorrencias.append({
            "processo_normalizado": normalizado,
            "valor_encontrado": correspondencia.group(0),
            "inicio": correspondencia.start(),
            "fim": correspondencia.end(),
        })

    return ocorrencias


def pesquisar_processos_em_registro(
    registro: dict,
    mapa_processos: dict[str, str],
    campos_pesquisa: Sequence[str] = CAMPOS_PESQUISA_PADRAO,
) -> list[dict]:
    """
    Retorna no máximo um resultado por processo dentro da publicação.

    Caso o processo apareça várias vezes ou em vários campos, todas as
    ocorrências são contabilizadas, mas o relatório mantém uma linha
    consolidada por processo/publicação.
    """
    processos_normalizados = set(mapa_processos)
    ocorrencias_por_processo: dict[str, list[dict]] = {
        processo: [] for processo in processos_normalizados
    }

    for campo in campos_pesquisa:
        texto_campo = texto_seguro(registro.get(campo))

        if not texto_campo:
            continue

        ocorrencias = localizar_processos_no_texto(
            texto=texto_campo,
            processos_normalizados=processos_normalizados,
        )

        for ocorrencia in ocorrencias:
            ocorrencia_completa = {
                **ocorrencia,
                "campo": campo,
                "texto_campo": texto_campo,
            }
            ocorrencias_por_processo[
                ocorrencia["processo_normalizado"]
            ].append(ocorrencia_completa)

    resultados: list[dict] = []

    for processo_normalizado, ocorrencias in ocorrencias_por_processo.items():
        if not ocorrencias:
            continue

        # Prioriza texto integral para o trecho principal.
        ocorrencias_ordenadas = sorted(
            ocorrencias,
            key=lambda item: (
                0 if item["campo"] == "texto_integral" else 1,
                campos_pesquisa.index(item["campo"])
                if item["campo"] in campos_pesquisa
                else 999,
                item["inicio"],
            ),
        )

        primeira = ocorrencias_ordenadas[0]
        campos_encontrados = list(dict.fromkeys(
            item["campo"] for item in ocorrencias_ordenadas
        ))

        processo_formatado = mapa_processos[processo_normalizado]
        id_publicacao = texto_seguro(
            primeiro_valor(
                registro,
                ["id_publicacao", "hash_conteudo", "indice_base"],
                "",
            )
        )

        chave_resultado = "|".join([
            processo_normalizado,
            id_publicacao,
            texto_seguro(registro.get("data_publicacao")),
            texto_seguro(registro.get("arquivo_xml")),
        ])

        resultados.append({
            "processo": processo_formatado,
            "processo_normalizado": processo_normalizado,
            "status": "ENCONTRADO",
            "data_publicacao": texto_seguro(
                primeiro_valor(
                    registro,
                    ["data_publicacao", "data_referencia", "data_xml"],
                    "",
                )
            ),
            "secao": texto_seguro(registro.get("secao")),
            "pagina": texto_seguro(registro.get("pagina")),
            "jornal": texto_seguro(registro.get("jornal")),
            "orgao": texto_seguro(registro.get("orgao")),
            "titulo": texto_seguro(
                primeiro_valor(
                    registro,
                    ["titulo", "titulo_publicacao", "titulo_listagem"],
                    "",
                )
            ),
            "ementa": texto_seguro(registro.get("ementa")),
            "identifica": texto_seguro(registro.get("identifica")),
            "sub_titulo": texto_seguro(registro.get("sub_titulo")),
            "quantidade_ocorrencias_publicacao": len(ocorrencias),
            "campos_encontrados": campos_encontrados,
            "campo_primeira_ocorrencia": primeira["campo"],
            "valor_encontrado": primeira["valor_encontrado"],
            "trecho_encontrado": montar_trecho(
                texto=primeira["texto_campo"],
                inicio=primeira["inicio"],
                fim=primeira["fim"],
            ),
            "url": texto_seguro(
                primeiro_valor(
                    registro,
                    [
                        "url",
                        "url_publicacao",
                        "link",
                        "href",
                        "url_consulta_dou",
                    ],
                    "",
                )
            ),
            "arquivo_xml": texto_seguro(registro.get("arquivo_xml")),
            "zip": texto_seguro(registro.get("zip")),
            "arquivo_origem": texto_seguro(registro.get("arquivo_origem")),
            "fonte": texto_seguro(registro.get("fonte")),
            "origem": texto_seguro(registro.get("origem")),
            "id_publicacao": texto_seguro(registro.get("id_publicacao")),
            "hash_conteudo": texto_seguro(registro.get("hash_conteudo")),
            "indice_base": registro.get("indice_base"),
            "chave_resultado": chave_resultado,
        })

    return resultados


def pesquisar_processos_nos_registros(
    registros: Iterable[dict],
    processos: Sequence[str] | None = None,
    campos_pesquisa: Sequence[str] = CAMPOS_PESQUISA_PADRAO,
) -> list[dict]:
    """
    Pesquisa todos os processos configurados em todos os registros da base.
    """
    mapa_processos = preparar_processos(processos)
    resultados: list[dict] = []
    chaves_vistas: set[str] = set()

    for registro in registros:
        if not isinstance(registro, dict):
            continue

        encontrados = pesquisar_processos_em_registro(
            registro=registro,
            mapa_processos=mapa_processos,
            campos_pesquisa=campos_pesquisa,
        )

        for resultado in encontrados:
            chave = resultado["chave_resultado"]

            if chave in chaves_vistas:
                continue

            chaves_vistas.add(chave)
            resultados.append(resultado)

    resultados.sort(
        key=lambda item: (
            item.get("data_publicacao") or "",
            item.get("processo") or "",
            item.get("orgao") or "",
            item.get("titulo") or "",
        )
    )

    return resultados


# ============================================================
# RESUMO
# ============================================================

def montar_resumo_processos(
    resultados: Sequence[dict],
    processos: Sequence[str] | None = None,
) -> dict:
    mapa_processos = preparar_processos(processos)

    publicacoes_por_processo: Counter[str] = Counter()
    ocorrencias_por_processo: Counter[str] = Counter()
    publicacoes_unicas: set[str] = set()

    for resultado in resultados:
        processo_normalizado = texto_seguro(
            resultado.get("processo_normalizado")
        )

        if processo_normalizado not in mapa_processos:
            continue

        publicacoes_por_processo[processo_normalizado] += 1
        ocorrencias_por_processo[processo_normalizado] += int(
            resultado.get("quantidade_ocorrencias_publicacao") or 0
        )

        identificador_publicacao = texto_seguro(
            resultado.get("id_publicacao")
            or resultado.get("hash_conteudo")
            or resultado.get("chave_resultado")
        )
        publicacoes_unicas.add(identificador_publicacao)

    itens = []

    for processo_normalizado, processo_formatado in mapa_processos.items():
        quantidade_publicacoes = publicacoes_por_processo[
            processo_normalizado
        ]
        quantidade_ocorrencias = ocorrencias_por_processo[
            processo_normalizado
        ]

        itens.append({
            "processo": processo_formatado,
            "processo_normalizado": processo_normalizado,
            "status": (
                "ENCONTRADO"
                if quantidade_publicacoes > 0
                else "NAO_ENCONTRADO"
            ),
            "quantidade_publicacoes": quantidade_publicacoes,
            "quantidade_ocorrencias": quantidade_ocorrencias,
        })

    return {
        "gerado_em": agora_iso(),
        "total_processos_pesquisados": len(mapa_processos),
        "total_processos_encontrados": sum(
            1 for item in itens if item["status"] == "ENCONTRADO"
        ),
        "total_processos_nao_encontrados": sum(
            1 for item in itens if item["status"] == "NAO_ENCONTRADO"
        ),
        "total_publicacoes_com_resultado": len(publicacoes_unicas),
        "total_linhas_resultado": len(resultados),
        "processos": itens,
    }


# ============================================================
# RELATÓRIOS
# ============================================================

def _serializar_valor_csv(valor: Any) -> Any:
    if isinstance(valor, (list, tuple, set)):
        return " | ".join(str(item) for item in valor)
    if isinstance(valor, dict):
        return json.dumps(valor, ensure_ascii=False)
    return valor


def salvar_csv_resultados(
    caminho: Path,
    resultados: Sequence[dict],
) -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)

    with caminho.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as arquivo:
        escritor = csv.DictWriter(
            arquivo,
            fieldnames=list(COLUNAS_CSV),
            delimiter=";",
            extrasaction="ignore",
        )
        escritor.writeheader()

        for resultado in resultados:
            linha = {
                coluna: _serializar_valor_csv(resultado.get(coluna, ""))
                for coluna in COLUNAS_CSV
            }
            escritor.writerow(linha)

    return caminho


def salvar_json_resultados(
    caminho: Path,
    resultados: Sequence[dict],
    resumo: dict,
    data_inicio: datetime.date | str,
    data_fim: datetime.date | str,
    metadados_execucao: dict | None = None,
) -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)

    payload = {
        "tipo_pesquisa": "PROCESSO_EXATO_SEM_IA",
        "data_inicio": str(data_inicio),
        "data_fim": str(data_fim),
        "gerado_em": agora_iso(),
        "resumo": resumo,
        "metadados_execucao": metadados_execucao or {},
        "resultados": list(resultados),
    }

    caminho.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return caminho


def salvar_txt_resumo(
    caminho: Path,
    resumo: dict,
    data_inicio: datetime.date | str,
    data_fim: datetime.date | str,
) -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)

    linhas = [
        "=" * 80,
        "PESQUISA EXATA DE PROCESSOS NO DOU — SEM IA",
        "=" * 80,
        f"Período: {data_inicio} a {data_fim}",
        f"Gerado em: {resumo.get('gerado_em')}",
        "-" * 80,
        (
            "Processos pesquisados: "
            f"{resumo.get('total_processos_pesquisados', 0)}"
        ),
        (
            "Processos encontrados: "
            f"{resumo.get('total_processos_encontrados', 0)}"
        ),
        (
            "Processos não encontrados: "
            f"{resumo.get('total_processos_nao_encontrados', 0)}"
        ),
        (
            "Publicações com resultado: "
            f"{resumo.get('total_publicacoes_com_resultado', 0)}"
        ),
        "-" * 80,
    ]

    for item in resumo.get("processos", []):
        linhas.append(
            f"{item.get('processo')} | "
            f"{item.get('status')} | "
            f"publicações={item.get('quantidade_publicacoes', 0)} | "
            f"ocorrências={item.get('quantidade_ocorrencias', 0)}"
        )

    linhas.append("=" * 80)

    caminho.write_text(
        "\n".join(linhas) + "\n",
        encoding="utf-8",
    )

    return caminho


def salvar_relatorios_pesquisa_processos(
    pasta_saida: Path | str,
    resultados: Sequence[dict],
    data_inicio: datetime.date | str,
    data_fim: datetime.date | str,
    processos: Sequence[str] | None = None,
    nome_base: str = "resultado_processos_2026",
    metadados_execucao: dict | None = None,
) -> dict:
    """
    Salva:
    - CSV detalhado;
    - JSON completo;
    - TXT com resumo.
    """
    pasta = Path(pasta_saida)
    pasta.mkdir(parents=True, exist_ok=True)

    resumo = montar_resumo_processos(
        resultados=resultados,
        processos=processos,
    )

    caminho_csv = salvar_csv_resultados(
        caminho=pasta / f"{nome_base}.csv",
        resultados=resultados,
    )

    caminho_json = salvar_json_resultados(
        caminho=pasta / f"{nome_base}.json",
        resultados=resultados,
        resumo=resumo,
        data_inicio=data_inicio,
        data_fim=data_fim,
        metadados_execucao=metadados_execucao,
    )

    caminho_txt = salvar_txt_resumo(
        caminho=pasta / f"{nome_base}_resumo.txt",
        resumo=resumo,
        data_inicio=data_inicio,
        data_fim=data_fim,
    )

    return {
        "csv": str(caminho_csv),
        "json": str(caminho_json),
        "txt": str(caminho_txt),
        "resumo": resumo,
    }


__all__ = [
    "PROCESSOS_PADRAO",
    "CAMPOS_PESQUISA_PADRAO",
    "normalizar_processo",
    "formatar_processo_normalizado",
    "carregar_processos_planilha",
    "preparar_processos",
    "localizar_processos_no_texto",
    "pesquisar_processos_em_registro",
    "pesquisar_processos_nos_registros",
    "montar_resumo_processos",
    "salvar_relatorios_pesquisa_processos",
]
